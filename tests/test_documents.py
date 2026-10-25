"""Extraction regression tests use generated fixtures, no patient documents."""

import io
import zipfile

import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject, NumberObject

from app.documents import extract_text, normalize_sections, parse_sections


def pdf_bytes(text: str = "", *, encrypt: bool = False, pages: int = 1) -> bytes:
    writer = PdfWriter()
    for _ in range(pages):
        page = writer.add_blank_page(width=612, height=792)
        if text:
            font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
            page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})})
            stream = DecodedStreamObject()
            stream.set_data(f"BT /F1 12 Tf 10 100 Td ({text}) Tj ET".encode("ascii"))
            page[NameObject("/Contents")] = writer._add_object(stream)
    if encrypt:
        writer.encrypt("password")
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


def test_txt_utf8_bom_and_gb18030_are_supported():
    sample = "【用法用量】\r\n模拟资料，仅用于测试。"
    expected = sample.replace("\r\n", "\n")
    assert extract_text("说明书.TXT", sample.encode("utf-8-sig")) == expected
    assert extract_text("说明书.txt", sample.encode("gb18030")) == expected


@pytest.mark.parametrize("filename,data", [
    ("empty.txt", b""),
    ("binary.txt", b"a\x00b"),
    ("old.doc", b"unsupported"),
    ("bad.pdf", b"not pdf"),
    ("bad.docx", b"not zip"),
    ("empty.txt", b" \n "),
    ("large.txt", b"x" * (10 * 1024 * 1024 + 1)),
])
def test_unsafe_or_unsupported_files_fail_cleanly(filename, data):
    with pytest.raises(ValueError):
        extract_text(filename, data)


def test_text_pdf_extracts_without_preserving_fake_page_numbers():
    result = extract_text("sample.pdf", pdf_bytes("SIMULATED LABEL 5 mg"))
    assert "SIMULATED LABEL 5 mg" in result


def test_blank_and_encrypted_pdf_have_actionable_errors():
    with pytest.raises(ValueError, match="OCR"):
        extract_text("blank.pdf", pdf_bytes())
    with pytest.raises(ValueError, match="加密"):
        extract_text("secret.pdf", pdf_bytes("example", encrypt=True))
    with pytest.raises(ValueError, match="100"):
        extract_text("many.pdf", pdf_bytes(pages=101))


def test_image_only_page_inside_mixed_pdf_is_not_silently_omitted():
    writer = PdfWriter()
    writer.append(io.BytesIO(pdf_bytes("TEXT ON FIRST PAGE")))
    page = writer.add_blank_page(width=100, height=100)
    image = DictionaryObject({NameObject("/Subtype"): NameObject("/Image")})
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/XObject"): DictionaryObject({NameObject("/Im0"): image})})
    output = io.BytesIO()
    writer.write(output)
    with pytest.raises(ValueError, match="第 2 页"):
        extract_text("mixed.pdf", output.getvalue())


def pdf_image_with_page_number(*, nested=False, inline=False) -> bytes:
    writer = PdfWriter()
    writer.append(io.BytesIO(pdf_bytes("Page 1")))
    page = writer.pages[0]
    content = page.get_contents().get_data()
    image_stream = DecodedStreamObject()
    image_stream.set_data(b"\x00\x00\x00")
    image_stream.update({NameObject("/Subtype"): NameObject("/Image"), NameObject("/Width"): NumberObject(1), NameObject("/Height"): NumberObject(1), NameObject("/ColorSpace"): NameObject("/DeviceRGB"), NameObject("/BitsPerComponent"): NumberObject(8)})
    if inline:
        content += b"\nq BI /W 1 /H 1 /BPC 8 /CS /RGB ID \x00\x00\x00 EI Q\n"
    else:
        image_reference = writer._add_object(image_stream)
        if nested:
            inner = DecodedStreamObject()
            inner.set_data(b"/Im0 Do")
            inner.update({NameObject("/Subtype"): NameObject("/Form"), NameObject("/Resources"): DictionaryObject({NameObject("/XObject"): DictionaryObject({NameObject("/Im0"): image_reference})})})
            outer = DecodedStreamObject()
            outer.set_data(b"/Inner Do")
            outer.update({NameObject("/Subtype"): NameObject("/Form"), NameObject("/Resources"): DictionaryObject({NameObject("/XObject"): DictionaryObject({NameObject("/Inner"): writer._add_object(inner)})})})
            image_reference = writer._add_object(outer)
        page["/Resources"][NameObject("/XObject")] = DictionaryObject({NameObject("/Visual"): image_reference})
        content += b"\nq /Visual Do Q\n"
    stream = DecodedStreamObject()
    stream.set_data(content)
    page[NameObject("/Contents")] = writer._add_object(stream)
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


@pytest.mark.parametrize("options", [{}, {"nested": True}, {"inline": True}])
def test_pdf_image_with_extractable_page_number_is_rejected(options):
    with pytest.raises(ValueError, match="此 PDF 含图片内容，可能遗漏文字"):
        extract_text("scan-with-page-number.pdf", pdf_image_with_page_number(**options))


def test_image_operator_names_inside_literal_text_are_allowed():
    text = "This is literal text: BI ID EI"
    assert text in extract_text("literal.pdf", pdf_bytes(text))


def test_docx_preserves_paragraph_table_paragraph_order():
    doc = Document()
    doc.add_paragraph("模拟示例，不用于临床")
    doc.add_paragraph("【药品名称】")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "测试样本 A"
    table.cell(0, 1).text = "模拟企业"
    doc.add_paragraph("【注意事项】全部数据为模拟。")
    output = io.BytesIO()
    doc.save(output)
    text = extract_text("test.docx", output.getvalue())
    assert text.index("药品名称") < text.index("测试样本 A") < text.index("模拟企业") < text.index("注意事项")
    assert parse_sections(text)[-1]["text"] == "全部数据为模拟。"


def test_docx_zip_bomb_rejected_before_xml_parse():
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", "x" * (2 * 1024 * 1024))
    with pytest.raises(ValueError, match="压缩比例"):
        extract_text("compressed.docx", output.getvalue())


def test_docx_header_and_footer_text_are_retained():
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "模拟页眉修订信息"
    doc.sections[0].footer.paragraphs[0].text = "模拟页脚重要文字"
    doc.add_paragraph("【药品名称】模拟品")
    output = io.BytesIO()
    doc.save(output)
    text = extract_text("headers.docx", output.getvalue())
    assert "模拟页眉修订信息" in text
    assert "模拟页脚重要文字" in text
    assert text.count("模拟页眉修订信息") == 1


def test_docx_unaccepted_revision_is_rejected():
    from docx.oxml import OxmlElement

    doc = Document()
    paragraph = doc.add_paragraph("模拟品")
    paragraph._p.append(OxmlElement("w:ins"))
    output = io.BytesIO()
    doc.save(output)
    with pytest.raises(ValueError, match="修订"):
        extract_text("revision.docx", output.getvalue())


def test_section_parser_preserves_preamble_aliases_duplicates_and_unknowns():
    sections = parse_sections("模拟示例，不用于临床\n【成分】内容一\n后续内容\n成份：内容二\n【企业自定义】附加说明")
    assert [s["key"] for s in sections] == ["前言", "成份", "成份#2", "企业自定义"]
    assert sections[1]["text"] == "内容一\n后续内容"
    assert (sections[1]["start_line"], sections[1]["end_line"]) == (2, 3)
    assert sections[2]["text"] == "内容二"


def test_unknown_unbracketed_text_is_never_dropped():
    text = "这是一段普通说明。\n[1] 是参考文献编号\n用法用量应由专业人员复核。"
    sections = parse_sections(text)
    assert len(sections) == 1
    assert sections[0]["key"] == "全文"
    assert sections[0]["text"] == text


def test_numbered_known_heading_and_heading_spaces():
    sections = parse_sections("一、用法与用量\n仅用于模拟\n[ 注 意 事 项 ]\n无临床用途")
    assert [s["key"] for s in sections] == ["用法用量", "注意事项"]


def test_manual_correction_clears_unverifiable_line_positions():
    sections = normalize_sections([
        {"title": "成分", "text": "模拟一", "start_line": 999},
        {"title": "成份", "text": "模拟二", "key": "untrusted"},
    ])
    assert [s["key"] for s in sections] == ["成份", "成份#2"]
    assert all(s["start_line"] is None and s["end_line"] is None for s in sections)


@pytest.mark.parametrize("sections", [[], [{"title": "", "text": "x"}], [{"title": "x", "text": ""}], [{"title": "x", "text": None}]])
def test_invalid_manual_corrections_rejected(sections):
    with pytest.raises(ValueError):
        normalize_sections(sections)
