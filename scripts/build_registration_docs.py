"""Build editable Word registration documents from verified project sources.

Uses python-docx only; PDF conversion is a separate LibreOffice step.
Source-code pagination changes only presentation, never the underlying files.
"""
from __future__ import annotations

import argparse
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import re
import subprocess
import unicodedata

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
BASE = ROOT / 'docs/registration'
SOURCE = BASE / 'source'
OUT = BASE / 'word'
NAME = '药品说明书版本比对与变更审核系统'
VERSION = 'V1.0'
COMMIT = '98e406f93f8777fc416014cee05d1737966fab4e'
SOURCE_ROWS_PER_PAGE = 60
FILES = ['run.py', 'app/__init__.py', 'app/main.py', 'app/database.py',
         'app/documents.py', 'app/comparison.py', 'app/static/styles.css',
         'app/static/index.html', 'app/static/app.js', 'app/reports.py', 'app/demo.py']
FONT = 'Noto Sans CJK SC'
MONO = 'Noto Sans Mono CJK SC'


def font(run, size=10.5, bold=False, color=None, mono=False):
    run.font.name = 'DejaVu Sans Mono' if mono else 'Arial'
    run.font.size = Pt(size)
    run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.insert(0, rfonts)
    rfonts.set(qn('w:eastAsia'), MONO if mono else FONT)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def field(paragraph, instruction):
    element = OxmlElement('w:fldSimple')
    element.set(qn('w:instr'), instruction)
    paragraph._p.append(element)


def setup(title, code=False):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.left_margin = section.right_margin = Cm(1.8 if code else 2.1)
    section.top_margin = section.bottom_margin = Cm(1.7 if code else 2.0)
    section.header_distance = section.footer_distance = Cm(.8)
    normal = doc.styles['Normal']
    normal.font.name = 'Arial'
    normal.font.size = Pt(10.5)
    normal._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), FONT)
    normal.paragraph_format.line_spacing = Pt(17)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.widow_control = True
    for name, size in [('Title', 24), ('Heading 1', 15), ('Heading 2', 12), ('Heading 3', 11)]:
        style = doc.styles[name]
        style.font.name = 'Arial'
        style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), FONT)
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string('173F3C')
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True
    header = section.header.paragraphs[0]
    font(header.add_run(f'{NAME} {VERSION}'), 8, color='526763')
    header.paragraph_format.space_after = Pt(0)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font(footer.add_run(f'{title}  |  第 '), 8, color='526763')
    field(footer, 'PAGE')
    font(footer.add_run(' 页 / 共 '), 8, color='526763')
    field(footer, 'NUMPAGES')
    font(footer.add_run(' 页'), 8, color='526763')
    doc.core_properties.title = f'{NAME} {VERSION} {title}'
    doc.core_properties.subject = '对应已实现软件的登记材料整理稿'
    doc.core_properties.author = '待申请人填写'
    doc.core_properties.keywords = 'MedLabel Review, V1.0'
    doc.core_properties.created = doc.core_properties.modified = datetime.now(timezone.utc)
    return doc


def soft_breaks(text):
    """Keep long file/test identifiers readable inside narrow table cells."""
    return re.sub(r'[A-Za-z][A-Za-z0-9_./]{34,}',
                  lambda match: re.sub(r'([_/\.])', lambda part: part.group(1) + '\u200b', match.group()),
                  text)


def inline(paragraph, text, size=10.5):
    tokens = re.split(r'(\*\*.*?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))', text)
    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**'):
            font(paragraph.add_run(token[2:-2]), size, bold=True)
        elif token.startswith('`') and token.endswith('`'):
            font(paragraph.add_run(soft_breaks(token[1:-1])), size - .5, mono=True)
        elif re.match(r'^\[[^\]]+\]\(', token):
            label, url = re.match(r'^\[([^\]]+)\]\(([^)]+)\)$', token).groups()
            font(paragraph.add_run(label), size, color='176B60')
            if url.startswith('http'):
                # A separate wrap-friendly URL avoids an unbroken run pushing
                # table content outside the page.
                font(paragraph.add_run('（' + url.replace('/', '/\u200b') + '）'), 8.5)
        else:
            font(paragraph.add_run(soft_breaks(token.replace('\\|', '|'))), size)


def cover(doc, title):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.keep_with_next = True
    font(p.add_run(f'{NAME} {VERSION}'), 16, bold=True, color='173F3C')
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_after = Pt(9)
    font(p.add_run(title), 15, bold=True, color='173F3C')
    for text in ['著作权人：待申请人填写', '对应代码版本：98e406f；文档状态：技术内容已编制，申请信息待确认']:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        font(p.add_run(text), 9, color='526763')


def add_table(doc, rows):
    if len(rows) < 2:
        return
    table = doc.add_table(rows=0, cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    for i, values in enumerate(rows):
        cells = table.add_row().cells
        if i == 0:
            repeat = OxmlElement('w:tblHeader')
            table.rows[-1]._tr.get_or_add_trPr().append(repeat)
        for j, cell in enumerate(cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_before = Pt(3)
            paragraph.paragraph_format.space_after = Pt(3)
            paragraph.paragraph_format.line_spacing = Pt(14)
            inline(paragraph, values[j] if j < len(values) else '', 9)
            if i == 0:
                shade = OxmlElement('w:shd')
                shade.set(qn('w:fill'), 'E7F0EC')
                cell._tc.get_or_add_tcPr().append(shade)
                for run in paragraph.runs:
                    run.bold = True
        nosplit = OxmlElement('w:cantSplit')
        table.rows[-1]._tr.get_or_add_trPr().append(nosplit)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def markdown_doc(path):
    title = path.stem.split('_', 1)[1]
    lines = path.read_text(encoding='utf-8').splitlines()
    doc = setup(title)
    cover(doc, title)
    i = 0
    while i < len(lines):
        line = lines[i]
        if not line.strip():
            i += 1
            continue
        if line.startswith('# ') or line.strip() == f'## {title}':
            i += 1
            continue
        if line.startswith('```'):
            block = []
            i += 1
            while i < len(lines) and not lines[i].startswith('```'):
                block.append(lines[i])
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = Pt(13)
            for j, item in enumerate(block):
                if j:
                    p.add_run().add_break()
                font(p.add_run(item), 8.5, mono=True)
            i += 1
            continue
        if line.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].startswith('|'):
                row = [v.strip() for v in re.split(r'(?<!\\)\|', lines[i].strip().strip('|'))]
                if not all(re.fullmatch(r':?-+:?', v.replace(' ', '')) for v in row):
                    rows.append(row)
                i += 1
            add_table(doc, rows)
            continue
        image = re.match(r'^!\[([^\]]*)\]\(([^)]+)\)', line)
        if image:
            target = (path.parent / image.group(2)).resolve()
            if not target.exists():
                raise FileNotFoundError(target)
            from PIL import Image
            with Image.open(target) as im:
                ratio = im.height / im.width
            width = min(16.7, 12.2 / ratio)
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            # Inline images need automatic line height; the body style's
            # exact 17 pt leading would otherwise crop the picture vertically.
            p.paragraph_format.line_spacing = 1
            p.paragraph_format.keep_with_next = True
            p.add_run().add_picture(str(target), width=Cm(width))
            caption = doc.add_paragraph()
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            caption.paragraph_format.space_after = Pt(9)
            font(caption.add_run(image.group(1)), 9, color='526763')
            i += 1
            continue
        heading = re.match(r'^(#{2,4})\s+(.+)', line)
        if heading:
            doc.add_heading(heading.group(2), level=len(heading.group(1)) - 1)
        else:
            p = doc.add_paragraph()
            if line.startswith('- '):
                line = '• ' + line[2:]
                p.paragraph_format.left_indent = Cm(.35)
                p.paragraph_format.first_line_indent = Cm(-.25)
            inline(p, line)
        i += 1
    OUT.mkdir(parents=True, exist_ok=True)
    output = OUT / (path.stem + '.docx')
    doc.save(output)
    return output


def wrap_source(text, width=82):
    text = text.expandtabs(4)
    pieces, part, used = [], '', 0
    for char in text:
        size = 2 if unicodedata.east_asian_width(char) in ('F', 'W') else 1
        if used + size > width:
            pieces.append(part)
            part, used = '', 0
        part += char
        used += size
    if part:
        pieces.append(part)
    return pieces


def source_material():
    rows, files, originals = [], [], []
    raw_count = nonempty_count = 0
    for filename in FILES:
        data = subprocess.check_output(['git', 'show', f'{COMMIT}:{filename}'], cwd=ROOT)
        if (ROOT / filename).read_bytes() != data:
            raise ValueError(f'Runtime source changed since fixed baseline: {filename}')
        text = data.decode('utf-8')
        original_lines = text.splitlines()
        record = {'path': filename, 'sha256': hashlib.sha256(data).hexdigest(),
                  'physical_lines': len(original_lines), 'nonempty_lines': sum(bool(s.strip()) for s in original_lines),
                  'first_display_row': len(rows) + 1}
        raw_count += record['physical_lines']
        nonempty_count += record['nonempty_lines']
        originals.append(f'===== 文件：{filename} =====\n{text}')
        for n, line in enumerate(original_lines, 1):
            if not line.strip():
                continue
            for continuation, segment in enumerate(wrap_source(line)):
                rows.append({'path': filename, 'source_line': n, 'continuation': continuation,
                             'text': segment, 'display_row': len(rows) + 1})
        record['last_display_row'] = len(rows)
        files.append(record)
    pages = [rows[i:i+SOURCE_ROWS_PER_PAGE] for i in range(0, len(rows), SOURCE_ROWS_PER_PAGE)]
    selected = list(range(len(pages))) if len(pages) <= 60 else list(range(30)) + list(range(len(pages)-30, len(pages)))
    doc = setup('源程序鉴别材料', code=True)
    source_map = []
    for submission_page, page_index in enumerate(selected, 1):
        if submission_page > 1:
            doc.add_page_break()
        items = pages[page_index]
        # This is page metadata, not an invented line of program text.
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = Pt(10)
        font(p.add_run(f'源文件范围：{items[0]["path"]}:{items[0]["source_line"]} → '
                       f'{items[-1]["path"]}:{items[-1]["source_line"]}'), 7.5)
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = Pt(10)
        font(p.add_run('左列为显示序号，次列为原始文件行号；+ 表示同一源行的续行。'), 7.5)
        for item in items:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            p.paragraph_format.line_spacing = Pt(10.5)
            p.paragraph_format.widow_control = False
            continuation = '+' if item['continuation'] else ' '
            prefix = f'{item["display_row"]:05d} {item["source_line"]:04d}{continuation} '
            font(p.add_run(prefix + item['text']), 8.5, mono=True)
        source_map.append({'submission_page': submission_page, 'full_layout_page': page_index + 1,
                           'first': {k:v for k,v in items[0].items() if k != 'text'},
                           'last': {k:v for k,v in items[-1].items() if k != 'text'}, 'display_rows': len(items),
                           'distinct_nonempty_source_lines': len({(x['path'], x['source_line']) for x in items})})
    OUT.mkdir(parents=True, exist_ok=True)
    doc.save(OUT / '05_源程序鉴别材料.docx')
    manifest = {'software': NAME, 'version': VERSION, 'baseline_commit': COMMIT,
                'files': files, 'total_physical_lines_including_blank': raw_count,
                'total_nonempty_source_lines': nonempty_count, 'total_display_rows': len(rows),
                'full_layout_pages': len(pages), 'submitted_pages': len(selected),
                'display_rows_per_full_page': SOURCE_ROWS_PER_PAGE,
                'page_policy': '每页60个显示行，末页50行；保留实际代码，省略纯空白行并对长行折行。全量75页，取前30页、后30页。显示行与原始源行分别标识，折行不计为新增源程序量。',
                'pages': source_map}
    (BASE / 'source_manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2)+'\n',encoding='utf-8')
    (BASE / '完整源程序汇编.txt').write_text('\n'.join(originals),encoding='utf-8')
    return manifest


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--source-only', action='store_true')
    args = parser.parse_args()
    manifest = source_material()
    print(json.dumps({k:manifest[k] for k in ['total_physical_lines_including_blank', 'total_nonempty_source_lines',
                                            'total_display_rows', 'full_layout_pages', 'submitted_pages']},ensure_ascii=False))
    if not args.source_only:
        for path in sorted(SOURCE.glob('*.md')):
            print(markdown_doc(path).relative_to(ROOT))


if __name__ == '__main__':
    main()
